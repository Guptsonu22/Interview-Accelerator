import io
import os
from typing import Dict, Any, Tuple
import pypdf
import docx
from backend.core.security import sanitize_untrusted_text, validate_file_upload, sanitize_filename


class DocumentParserService:
    @staticmethod
    def parse_pdf_bytes(file_bytes: bytes) -> Tuple[str, Dict[str, Any]]:
        """Extract text and metadata from PDF bytes."""
        try:
            reader = pypdf.PdfReader(io.BytesIO(file_bytes))
            if reader.is_encrypted:
                try:
                    reader.decrypt("")
                except Exception:
                    raise ValueError("PDF is password protected and cannot be extracted.")

            text_parts = []
            for i, page in enumerate(reader.pages):
                extracted = page.extract_text()
                if extracted:
                    text_parts.append(extracted)

            combined_text = "\n\n".join(text_parts).strip()
            metadata = {
                "format": "pdf",
                "page_count": len(reader.pages),
                "extracted_chars": len(combined_text)
            }
            return combined_text, metadata
        except Exception as e:
            if "password" in str(e).lower():
                raise ValueError("PDF is encrypted and password-protected.")
            raise ValueError(f"Failed to parse PDF document: {str(e)}")

    @staticmethod
    def parse_docx_bytes(file_bytes: bytes) -> Tuple[str, Dict[str, Any]]:
        """Extract text and metadata from DOCX bytes."""
        try:
            doc = docx.Document(io.BytesIO(file_bytes))
            text_parts = []
            
            # Paragraphs
            for p in doc.paragraphs:
                clean_p = p.text.strip()
                if clean_p:
                    text_parts.append(clean_p)
                    
            # Tables
            for table in doc.tables:
                for row in table.rows:
                    row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_texts:
                        text_parts.append(" | ".join(row_texts))

            combined_text = "\n\n".join(text_parts).strip()
            metadata = {
                "format": "docx",
                "paragraph_count": len(doc.paragraphs),
                "table_count": len(doc.tables),
                "extracted_chars": len(combined_text)
            }
            return combined_text, metadata
        except Exception as e:
            raise ValueError(f"Failed to parse DOCX document: {str(e)}")

    @staticmethod
    def parse_text_bytes(file_bytes: bytes) -> Tuple[str, Dict[str, Any]]:
        """Decode plain text bytes supporting UTF-8 and latin-1 fallbacks."""
        for enc in ["utf-8", "utf-8-sig", "latin-1", "cp1252"]:
            try:
                decoded = file_bytes.decode(enc)
                return decoded.strip(), {"format": "txt", "encoding": enc, "extracted_chars": len(decoded)}
            except UnicodeDecodeError:
                continue
        raise ValueError("Could not decode text file using supported encodings.")

    @classmethod
    def parse_file(cls, filename: str, file_bytes: bytes) -> Tuple[str, Dict[str, Any]]:
        """Validate and parse an uploaded file (PDF, DOCX, or TXT)."""
        clean_name = sanitize_filename(filename)
        valid, err = validate_file_upload(clean_name, len(file_bytes))
        if not valid:
            raise ValueError(err)

        _, ext = os.path.splitext(clean_name.lower())
        if ext == ".pdf":
            raw_text, meta = cls.parse_pdf_bytes(file_bytes)
        elif ext == ".docx":
            raw_text, meta = cls.parse_docx_bytes(file_bytes)
        elif ext in [".txt", ".text"]:
            raw_text, meta = cls.parse_text_bytes(file_bytes)
        else:
            raise ValueError(f"Unsupported file extension: {ext}")

        sanitized = sanitize_untrusted_text(raw_text)
        if len(sanitized.strip()) < 20:
            raise ValueError("Extracted document contains insufficient text (fewer than 20 characters). Please provide a valid document.")

        meta["sanitized_chars"] = len(sanitized)
        meta["filename"] = clean_name
        return sanitized, meta
